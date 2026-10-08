import os
import datetime
import json
import logging
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..database import SessionLocal
from ..models import Job, SchoolRecord

logger = logging.getLogger("schoolminer.export")

FIELD_MAP = {
    "School Name": "school_name",
    "UDISE": "udise",
    "Affiliation Number": "affiliation_number",
    "Address": "address",
    "Village": "village",
    "City": "city",
    "District": "district",
    "State": "state",
    "PIN Code": "pin_code",
    "Phone": "phone",
    "Mobile": "mobile",
    "Email": "email",
    "Website": "website",
    "Principal": "principal",
    "Category": "category",
    "Management": "management",
    "School Type": "school_type",
    "Medium": "medium",
    "Latitude": "latitude",
    "Longitude": "longitude",
    "Established Year": "established_year"
}

def generate_export_files(job_id: int):
    """Generate Excel (xlsx) and CSV reports for the scraped school records."""
    db = SessionLocal()
    try:
        job = db.query(Job).filter(Job.id == job_id).first()
        if not job:
            logger.error(f"Job {job_id} not found for export.")
            return
            
        records = db.query(SchoolRecord).filter(SchoolRecord.job_id == job_id).all()
        if not records:
            logger.warning(f"No records found for job {job_id}. Skipping export.")
            return
            
        # Parse selected fields
        try:
            selected_fields = json.loads(job.fields)
        except Exception:
            selected_fields = []
            
        if not selected_fields:
            selected_fields = list(FIELD_MAP.keys())
            
        # Create output directory
        out_dir = job.output_folder or "Output"
        os.makedirs(out_dir, exist_ok=True)
        
        # Build clean filenames
        clean_state = job.state.replace(" ", "_").lower()
        clean_districts = job.districts.replace(" ", "_").replace(",", "_").lower()
        # Keep name clean and reasonably short
        if len(clean_districts) > 50:
            clean_districts = clean_districts[:47] + "_etc"
        filename_base = f"schoolminer_{job_id}_{clean_state}_{clean_districts}"
        
        excel_path = os.path.join(out_dir, f"{filename_base}.xlsx")
        csv_path = os.path.join(out_dir, f"{filename_base}.csv")
        
        # Prepare list of dictionaries matching user selected fields
        records_data = []
        for r in records:
            row = {}
            for field in selected_fields:
                db_col = FIELD_MAP.get(field)
                if db_col:
                    row[field] = getattr(r, db_col, "") or ""
            records_data.append(row)
            
        df = pd.DataFrame(records_data)
        
        # Export CSV if requested
        if job.output_format.lower() in ("csv", "both"):
            df.to_csv(csv_path, index=False, encoding="utf-8")
            logger.info(f"CSV report written: {csv_path}")
            
        # Export styled Excel if requested
        if job.output_format.lower() in ("xlsx", "both"):
            wb = Workbook()
            
            # --- Sheet 1: Summary Sheet ---
            ws_summary = wb.active
            ws_summary.title = "Summary"
            ws_summary.views.sheetView[0].showGridLines = True
            
            # Styles
            title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
            title_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
            bold_font = Font(name="Calibri", size=11, bold=True)
            regular_font = Font(name="Calibri", size=11)
            header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )
            
            # Write Title Banner
            ws_summary.merge_cells("A1:C1")
            title_cell = ws_summary["A1"]
            title_cell.value = "SchoolMiner Enterprise Scraper Summary"
            title_cell.font = title_font
            title_cell.fill = title_fill
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws_summary.row_dimensions[1].height = 40
            
            # Write Metdata Tables
            summary_info = [
                ("Export Parameter", "Value", ""),  # Header
                ("Job ID", job.id, "Unique run identifier"),
                ("Export Timestamp", datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"), "Report generation time"),
                ("Export Board", job.board.upper(), "Target board directory"),
                ("Target State", job.state.title(), "Target geographical state"),
                ("Scraped Districts", job.districts, "Scope of target districts"),
                ("Total Records Exported", len(records), "Count of deduplicated schools"),
                ("Scrape Speed", f"{job.speed_rpm} RPM", "Records scraped per minute"),
                ("Errors Encountered", job.errors_count, "Errors resolved by retry scheduler"),
                ("Job Final Status", job.status, "Execution ending state"),
                ("Software Version", "v5.0-Enterprise", "Application edition code")
            ]
            
            for row_idx, (param, val, desc) in enumerate(summary_info, start=3):
                ws_summary.cell(row=row_idx, column=1, value=param)
                ws_summary.cell(row=row_idx, column=2, value=val)
                ws_summary.cell(row=row_idx, column=3, value=desc)
                
                # Styles
                c1 = ws_summary.cell(row=row_idx, column=1)
                c2 = ws_summary.cell(row=row_idx, column=2)
                c3 = ws_summary.cell(row=row_idx, column=3)
                
                c1.border = thin_border
                c2.border = thin_border
                c3.border = thin_border
                
                if row_idx == 3:  # Table header
                    c1.font = bold_font
                    c2.font = bold_font
                    c3.font = bold_font
                    c1.fill = header_fill
                    c2.fill = header_fill
                    c3.fill = header_fill
                else:
                    c1.font = bold_font
                    c2.font = regular_font
                    c3.font = regular_font
                    
            ws_summary.column_dimensions["A"].width = 25
            ws_summary.column_dimensions["B"].width = 30
            ws_summary.column_dimensions["C"].width = 35
            
            # --- Sheet 2: Data Sheet ---
            ws_data = wb.create_sheet(title="Schools Data")
            ws_data.views.sheetView[0].showGridLines = True
            
            # Write Header Row
            headers = df.columns.tolist()
            ws_data.append(headers)
            
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            data_header_fill = PatternFill(start_color="2D3748", end_color="2D3748", fill_type="solid")
            center_align = Alignment(horizontal="center", vertical="center")
            
            for col_idx, header in enumerate(headers, start=1):
                cell = ws_data.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = data_header_fill
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
            ws_data.row_dimensions[1].height = 28
            
            # Write Data Rows
            for r_idx, row_data in enumerate(records_data, start=2):
                ws_data.append([row_data[h] for h in headers])
                ws_data.row_dimensions[r_idx].height = 20
                for c_idx in range(1, len(headers) + 1):
                    cell = ws_data.cell(row=r_idx, column=c_idx)
                    cell.font = regular_font
                    cell.border = thin_border
                    # Align numbers / coordinates
                    if headers[c_idx-1] in ("PIN Code", "Phone", "Mobile", "Latitude", "Longitude", "Established Year", "UDISE", "Affiliation Number"):
                        cell.alignment = Alignment(horizontal="center")
            
            # Enable Auto Filters
            max_col_letter = get_column_letter(len(headers))
            ws_data.auto_filter.ref = f"A1:{max_col_letter}{len(records_data) + 1}"
            
            # Freeze First Row
            ws_data.freeze_panes = "A2"
            
            # Auto-adjust column widths
            for col in ws_data.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws_data.column_dimensions[col_letter].width = max(max_len + 3, 12)
                
            wb.save(excel_path)
            logger.info(f"Excel report written: {excel_path}")
            
    except Exception as e:
        logger.error(f"Failed to generate reports for job {job_id}: {e}")
    finally:
        db.close()
